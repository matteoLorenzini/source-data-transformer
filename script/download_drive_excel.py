"""Download Google Sheets and transform worksheet data with external XSLT.

Usage:
    download_drive_excel.py --file <LOCAL_XLSX_PATH>
    download_drive_excel.py --file-id <GOOGLE_DRIVE_FILE_ID>
    download_drive_excel.py --url <GOOGLE_SHEETS_SHARE_URL>
    download_drive_excel.py --url <GOOGLE_SHEETS_SHARE_URL> --stylesheet-map <JSON_MAP>

The script serializes raw worksheet data, then applies the default stylesheet
or one selected by spreadsheet file ID to produce CSV, XML, and relation XML.
"""

from __future__ import annotations

import argparse
import csv
import json
import os
import re
from pathlib import Path
from typing import Iterable
from urllib.parse import unquote
from xml.etree.ElementTree import Element, SubElement, tostring

import requests
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_OUTPUT_DIR = SCRIPT_DIR / "output_data"


def print_credentials_help() -> None:
    """Print the supported authentication options for private sheets."""

    print("Credential options for private Google Sheets:")
    print("1. Service account JSON")
    print("   - Create a service account in Google Cloud.")
    print("   - Download the JSON key file.")
    print("   - Share the target Google Sheet with the service account email.")
    print("   - Run: python download_drive_excel.py --url <SHEET_URL> --service-account-json <PATH_TO_JSON>")
    print("2. OAuth2 access token")
    print("   - Generate a bearer token with access to the target Google Sheet.")
    print("   - Run: python download_drive_excel.py --url <SHEET_URL> --access-token <TOKEN>")
    print("   - Or set GOOGLE_ACCESS_TOKEN in the environment.")
    print("3. Public sheet")
    print("   - Make the Google Sheet readable by anyone with the link.")
    print("   - Run the script without authentication arguments.")


def extract_file_id(value: str) -> str:
    """Extract a Google Drive file id from a raw id or share URL."""

    patterns = [
        r"/d/([a-zA-Z0-9_-]+)",
        r"id=([a-zA-Z0-9_-]+)",
        r"/file/d/([a-zA-Z0-9_-]+)",
    ]
    for pattern in patterns:
        match = re.search(pattern, value)
        if match:
            return match.group(1)
    return value.strip()


def extract_sheet_gid(value: str) -> str | None:
    """Extract a Google Sheets gid from a URL when one is present."""

    match = re.search(r"[?#&]gid=(\d+)", value)
    if match:
        return match.group(1)
    return None


def sanitize_filename(value: str) -> str:
    """Turn a worksheet title into a safe file name."""

    safe_value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return safe_value.strip("._") or "sheet"


def load_stylesheet_mapping(mapping_path: Path) -> dict[str, Path]:
    """Load Google Drive file-id to XSLT path mappings from JSON."""

    try:
        configuration = json.loads(mapping_path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid stylesheet mapping JSON in {mapping_path}: {exc}") from exc

    if not isinstance(configuration, dict):
        raise ValueError("The stylesheet mapping must be a JSON object of file IDs to stylesheet paths.")

    mapping: dict[str, Path] = {}
    for file_id, stylesheet in configuration.items():
        if not isinstance(file_id, str) or not isinstance(stylesheet, str) or not stylesheet.strip():
            raise ValueError("Every stylesheet mapping must have a file ID and a non-empty path.")

        stylesheet_path = Path(stylesheet)
        if not stylesheet_path.is_absolute():
            stylesheet_path = mapping_path.parent / stylesheet_path
        mapping[file_id] = stylesheet_path.resolve()

    return mapping


def load_stylesheet(stylesheet_path: Path):
    """Compile an external XSLT stylesheet."""

    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    try:
        stylesheet_tree = etree.parse(str(stylesheet_path), parser)
        return etree.XSLT(
            stylesheet_tree,
            access_control=etree.XSLTAccessControl.DENY_ALL,
        )
    except (etree.XMLSyntaxError, etree.XSLTParseError) as exc:
        raise ValueError(f"Could not load XSLT stylesheet {stylesheet_path}: {exc}") from exc


def apply_stylesheet(source_root: Element, stylesheet, output_mode: str):
    """Transform neutral source XML using the selected output mode."""

    from lxml import etree

    source_tree = etree.fromstring(tostring(source_root, encoding="utf-8"))
    try:
        return stylesheet(source_tree, output_mode=etree.XSLT.strparam(output_mode))
    except etree.XSLTApplyError as exc:
        raise ValueError(f"XSLT transformation failed: {exc}") from exc


def write_xslt_xml(result, destination_path: Path) -> None:
    """Write an XSLT XML result, defaulting an unspecified encoding to UTF-8."""

    try:
        result.write_output(str(destination_path))
    except LookupError as exc:
        if str(exc) != "unknown encoding: ''":
            raise
        from lxml import etree

        destination_path.write_bytes(etree.tostring(result, encoding="utf-8", xml_declaration=True))


def source_cell_text(value) -> str:
    """Convert a spreadsheet value to source text without applying field rules."""

    return "" if value is None else str(value)


def extract_sheet_name_from_content_disposition(
    content_disposition: str | None,
    fallback: str,
) -> str:
    """Extract sheet title from Content-Disposition header when available."""

    if not content_disposition:
        return fallback

    filename_star_match = re.search(r"filename\*=UTF-8''([^;]+)", content_disposition, re.IGNORECASE)
    if filename_star_match:
        value = unquote(filename_star_match.group(1)).strip().strip('"')
    else:
        filename_match = re.search(r"filename=\"?([^\";]+)\"?", content_disposition, re.IGNORECASE)
        if not filename_match:
            return fallback
        value = filename_match.group(1).strip().strip('"')

    if value.lower().endswith(".csv"):
        value = value[:-4]

    return value.strip() or fallback


def sanitize_xml_tag(value: str, fallback: str) -> str:
    """Convert a header value into a safe XML element name."""

    normalized = source_cell_text(value).strip()
    tag = re.sub(r"[^A-Za-z0-9_.-]+", "_", normalized)
    tag = re.sub(r"_+", "_", tag).strip("_.-")

    if not tag:
        tag = fallback

    if not re.match(r"^[A-Za-z_]", tag) or tag.lower().startswith("xml"):
        tag = f"col_{tag}"

    return tag


def canonical_header_key(value: str) -> str:
    """Normalize a header name for stylesheet rule matching."""

    return sanitize_xml_tag(value, "col").upper()


def build_column_tags(header_row: list[str], max_columns: int) -> list[str]:
    """Build unique XML element names from the worksheet header row."""

    tags: list[str] = []
    seen: dict[str, int] = {}

    for column_index in range(1, max_columns + 1):
        header_value = header_row[column_index - 1] if column_index - 1 < len(header_row) else ""
        fallback = f"col_{get_column_letter(column_index)}"
        base_tag = sanitize_xml_tag(header_value, fallback)

        seen_count = seen.get(base_tag, 0)
        seen[base_tag] = seen_count + 1
        tag_name = base_tag if seen_count == 0 else f"{base_tag}_{seen_count + 1}"
        tags.append(tag_name)

    return tags


def build_http_session() -> requests.Session:
    """Create an HTTP session with browser-like headers for Google exports."""

    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/126.0 Safari/537.36",
            "Accept": "text/csv,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/octet-stream;q=0.9,*/*;q=0.8",
            "Accept-Language": "en-US,en;q=0.9",
        }
    )
    return session


def resolve_access_token(
    cli_access_token: str | None,
    service_account_json: str | None,
) -> str | None:
    """Resolve an OAuth2 bearer token from CLI, env var, or service account."""

    if cli_access_token:
        return cli_access_token

    env_token = os.getenv("GOOGLE_ACCESS_TOKEN")
    if env_token:
        return env_token

    if not service_account_json:
        return None

    try:
        from google.auth.transport.requests import Request
        from google.oauth2 import service_account
    except ImportError as exc:
        raise RuntimeError(
            "Service account auth requires google-auth. Install it with: pip install google-auth"
        ) from exc

    credentials = service_account.Credentials.from_service_account_file(
        service_account_json,
        scopes=[
            "https://www.googleapis.com/auth/spreadsheets.readonly",
            "https://www.googleapis.com/auth/drive.readonly",
        ],
    )
    credentials.refresh(Request())
    return credentials.token


def raise_export_error(response: requests.Response, file_id: str, gid: str | None = None) -> None:
    """Raise a helpful error for Google export permission failures."""

    if response.status_code in (401, 403):
        target = f"file {file_id}" if gid is None else f"file {file_id} gid {gid}"
        raise RuntimeError(
            "Google denied access to "
            f"{target}. Make the sheet publicly readable or provide "
            "--access-token (or GOOGLE_ACCESS_TOKEN) for an authenticated export."
        )
    response.raise_for_status()


def download_google_sheet_as_xlsx(file_id: str, destination: Path, access_token: str | None = None) -> None:
    """Download a Google Sheets document as an .xlsx export."""

    url = f"https://docs.google.com/spreadsheets/d/{file_id}/export"
    session = build_http_session()
    if access_token:
        session.headers["Authorization"] = f"Bearer {access_token}"
    response = session.get(url, params={"format": "xlsx"}, stream=True, timeout=60)

    raise_export_error(response, file_id)

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("wb") as output_file:
        for chunk in response.iter_content(chunk_size=1024 * 1024):
            if chunk:
                output_file.write(chunk)


def download_google_sheet_as_csv(
    file_id: str,
    gid: str,
    output_dir: Path,
    access_token: str | None = None,
) -> tuple[Path, str]:
    """Download a single Google Sheets worksheet as CSV and return its path and sheet name."""

    url = f"https://docs.google.com/spreadsheets/d/{file_id}/export"
    session = build_http_session()
    if access_token:
        session.headers["Authorization"] = f"Bearer {access_token}"
    response = session.get(url, params={"format": "csv", "gid": gid}, stream=True, timeout=60)

    raise_export_error(response, file_id, gid)

    sheet_name = extract_sheet_name_from_content_disposition(
        response.headers.get("Content-Disposition"),
        fallback=f"gid_{gid}",
    )
    sheet_slug = sanitize_filename(sheet_name)
    destination = output_dir / f"{sheet_slug}.csv"

    destination.parent.mkdir(parents=True, exist_ok=True)
    with destination.open("wb") as output_file:
        for chunk in response.iter_content(chunk_size=1024 * 1024):
            if chunk:
                output_file.write(chunk)

    return destination, sheet_name


def rows_to_source_xml(rows: list[list]) -> Element:
    """Serialize raw tabular values and header metadata for XSLT processing."""

    header_row = [source_cell_text(value) for value in rows[0]] if rows else []
    max_columns = max(len(header_row), max((len(row) for row in rows), default=0))
    column_tags = build_column_tags(header_row, max_columns)

    source_root = Element("source")
    columns_element = SubElement(source_root, "columns")
    for column_index, tag_name in enumerate(column_tags):
        header_value = header_row[column_index] if column_index < len(header_row) else ""
        column_element = SubElement(
            columns_element,
            "column",
            index=str(column_index),
            key=canonical_header_key(header_value),
            tag=tag_name,
        )
        column_element.text = header_value

    rows_element = SubElement(source_root, "rows")
    for row_index, row in enumerate(rows[1:], start=2):
        row_element = SubElement(rows_element, "row", index=str(row_index))
        for column_index, tag_name in enumerate(column_tags):
            value = row[column_index] if column_index < len(row) else ""
            cell_element = SubElement(
                row_element,
                "cell",
                index=str(column_index),
                key=canonical_header_key(header_row[column_index] if column_index < len(header_row) else ""),
                tag=tag_name,
            )
            cell_element.text = source_cell_text(value)

    return source_root


def worksheet_to_source_xml(worksheet) -> Element:
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    while rows and not any(value is not None for value in rows[-1]):
        rows.pop()
    return rows_to_source_xml(rows)


def write_source_exports(
    source_root: Element,
    csv_path: Path,
    xml_path: Path,
    stylesheet,
) -> None:
    """Generate all exported formats through the selected external stylesheet."""

    csv_path.parent.mkdir(parents=True, exist_ok=True)
    xml_path.parent.mkdir(parents=True, exist_ok=True)

    csv_result = apply_stylesheet(source_root, stylesheet, "csv")
    with csv_path.open("w", newline="", encoding="utf-8") as csv_file:
        writer = csv.writer(csv_file)
        for row_element in csv_result.getroot().findall("row"):
            writer.writerow([
                cell_element.text or ""
                for cell_element in row_element.findall("cell")
            ])

    main_result = apply_stylesheet(source_root, stylesheet, "main")
    write_xslt_xml(main_result, xml_path)

    relation_xml_path = xml_path.with_name(f"{xml_path.stem}_relations{xml_path.suffix}")
    relation_result = apply_stylesheet(source_root, stylesheet, "relations")
    write_xslt_xml(relation_result, relation_xml_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Download a Google Sheet and export each worksheet to CSV and XML.")
    parser.add_argument(
        "--credentials-help",
        action="store_true",
        help="Print authentication setup instructions for private Google Sheets and exit.",
    )
    parser.add_argument(
        "--file",
        type=Path,
        help="Local .xlsx workbook to process without downloading from Google.",
    )
    parser.add_argument(
        "--url",
        help="Google Drive share URL containing the spreadsheet file id.",
    )
    parser.add_argument(
        "--file-id",
        help="Google Drive file id. Use this if you already know the id.",
    )
    parser.add_argument(
        "--gid",
        help="Google Sheets gid. If omitted, the script uses gid from the URL when present.",
    )
    parser.add_argument(
        "--sheet-name",
        help="Export only this worksheet from a workbook; by default all worksheets are exported.",
    )
    parser.add_argument(
        "--access-token",
        help="OAuth2 bearer token for private Google Sheets export. Can also be set via GOOGLE_ACCESS_TOKEN.",
    )
    parser.add_argument(
        "--service-account-json",
        help="Path to a Google service account key JSON used to generate an access token.",
    )
    parser.add_argument(
        "--output-dir",
        default=str(DEFAULT_OUTPUT_DIR),
        help="Directory where the downloaded file and exports will be written.",
    )
    parser.add_argument(
        "--base-name",
        default="google_sheet_export",
        help="Base folder name used for the downloaded Excel file and worksheet exports.",
    )
    parser.add_argument(
        "--stylesheet-map",
        type=Path,
        help="XSLT file to use directly, or JSON mapping of Google Drive file IDs to XSLT paths.",
    )
    return parser


def main(argv: Iterable[str] | None = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)

    if args.credentials_help:
        print_credentials_help()
        return 0

    source_count = sum(bool(value) for value in (args.file, args.url, args.file_id))
    if source_count != 1:
        parser.error("provide exactly one of --file, --url, or --file-id")

    local_file_path = args.file.resolve() if args.file else None
    if local_file_path is not None:
        if not local_file_path.is_file():
            parser.error(f"Local workbook does not exist: {local_file_path}")
        if local_file_path.suffix.lower() != ".xlsx":
            parser.error("--file currently supports .xlsx workbooks")
        if args.gid:
            parser.error("--gid can only be used with --url or --file-id")
        stylesheet_key = local_file_path.stem
        file_id = None
        gid = None
        access_token = None
    else:
        raw_id = args.file_id or args.url
        assert raw_id is not None
        file_id = extract_file_id(raw_id)
        stylesheet_key = file_id
        gid = args.gid or extract_sheet_gid(raw_id)
        access_token = resolve_access_token(args.access_token, args.service_account_json)

    if gid is not None and args.sheet_name:
        parser.error("--sheet-name cannot be used with --gid, which already selects a single worksheet")

    stylesheet_path = SCRIPT_DIR / "xslt" / "default.xsl"
    if args.stylesheet_map:
        if args.stylesheet_map.suffix.lower() in {".xsl", ".xslt"}:
            stylesheet_path = args.stylesheet_map.resolve()
        else:
            try:
                stylesheet_mapping = load_stylesheet_mapping(args.stylesheet_map)
            except (OSError, ValueError) as exc:
                parser.error(str(exc))

            stylesheet_path = stylesheet_mapping.get(stylesheet_key)
            if stylesheet_path is None:
                parser.error(f"No XSLT stylesheet configured for source {stylesheet_key}.")

    if not stylesheet_path.is_file():
        parser.error(f"XSLT stylesheet does not exist: {stylesheet_path}")
    try:
        stylesheet = load_stylesheet(stylesheet_path)
    except (OSError, ValueError) as exc:
        parser.error(str(exc))

    output_dir = Path(args.output_dir)
    workbook_dir = output_dir / args.base_name
    excel_path = workbook_dir / f"{args.base_name}.xlsx"

    if local_file_path is not None:
        print(f"Reading local workbook from {local_file_path}...")
        workbook = load_workbook(local_file_path, data_only=True)
    elif gid is not None:
        print(f"Downloading Google Sheet {file_id} gid {gid} as CSV...")
        csv_path, sheet_name = download_google_sheet_as_csv(
            file_id,
            gid,
            workbook_dir,
            access_token=access_token,
        )

        sheet_slug = sanitize_filename(sheet_name)
        xml_path = workbook_dir / f"{sheet_slug}.xml"

        print(f"Reading CSV from {csv_path}...")
        with csv_path.open("r", encoding="utf-8", newline="") as csv_file:
            reader = csv.reader(csv_file)
            rows = [row for row in reader]

        print(f"Writing XML to {xml_path}...")
        source_root = rows_to_source_xml(rows)
        try:
            write_source_exports(source_root, csv_path, xml_path, stylesheet)
        except ValueError as exc:
            parser.error(str(exc))
    else:
        print(f"Downloading Google Sheet {file_id} as Excel...")
        download_google_sheet_as_xlsx(file_id, excel_path, access_token=access_token)

        print(f"Reading workbook from {excel_path}...")
        workbook = load_workbook(excel_path, data_only=True)

    if gid is None:
        if args.sheet_name:
            if args.sheet_name not in workbook.sheetnames:
                parser.error(
                    f"Worksheet {args.sheet_name!r} not found. Available worksheets: "
                    f"{', '.join(workbook.sheetnames)}"
                )
            worksheets = [workbook[args.sheet_name]]
        else:
            worksheets = workbook.worksheets

        for worksheet in worksheets:
            sheet_slug = sanitize_filename(worksheet.title)
            csv_path = workbook_dir / f"{sheet_slug}.csv"
            xml_path = workbook_dir / f"{sheet_slug}.xml"

            print(f"Writing CSV to {csv_path}...")
            print(f"Writing XML to {xml_path}...")
            source_root = worksheet_to_source_xml(worksheet)
            try:
                write_source_exports(source_root, csv_path, xml_path, stylesheet)
            except ValueError as exc:
                parser.error(str(exc))

    print("Done.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())